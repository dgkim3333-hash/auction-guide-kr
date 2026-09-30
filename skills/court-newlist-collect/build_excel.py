# -*- coding: utf-8 -*-
"""
court-newlist-collect / build_excel.py
build_rows.py 가 만든 CSV → 02단원 3절 규격의 신건레이더 엑셀 (탐색 시트).

  1단계  python build_excel.py stage1 court_newlist_YYYYMMDD.csv --out 신건레이더_YYYYMMDD.xlsx [--title "..."]
         → 값·헤더·서식·행 높이·열 폭·보조열 수식·틀고정을 openpyxl 로 쓴다. 표·수식·슬라이서는 아직 없다.
         → 끝에 2단계에서 Excel MCP 에 넣을 값(표 범위·수식·슬라이서 목록)을 JSON 으로 출력한다.
  2단계  Excel MCP (SKILL.md STEP 4) — 표 생성 → C4:C9 수식 → A15 스필 수식 → 슬라이서 8개 → 저장·닫기
  3단계  python build_excel.py patch 신건레이더_YYYYMMDD.xlsx
         → 슬라이서를 144×170pt 절대 좌표로, 머리글을 열 이름으로 바꾸고 검사한다. 검사 실패면 원본을 건드리지 않는다.

규격 출처: 02단원 3절 [3]·[4] (2026-09-30 Cowork·Code 실측으로 두 환경의 산출물이 달라져 스크립트로 고정)
금액은 원 단위 정수. 표 이름은 영문 TBL_new (Excel MCP 가 한글 표 이름을 못 만든다).
"""
import argparse, csv, json, os, re, sys, zipfile
from datetime import date

FONT = "맑은 고딕"
TOPHDR = 14
TABLE = "TBL_new"
# (열이름, 폭) — build_rows.py 출력 19열 + 연번
HEAD = [('연번', 6), ('사건번호', 14), ('법원', 10), ('매각기일', 12), ('소재지', 44), ('용도', 14),
        ('감정가', 16), ('최저가', 16), ('최저가율', 9), ('건물면적', 10), ('토지면적', 10), ('조회수', 8),
        ('비고', 18), ('시도', 7), ('시군구', 10), ('용도군', 10), ('감정가 규모대', 28), ('남은 날', 12),
        ('조회수 구간', 22), ('신규', 7)]
SLICERS = [('SL_A', '용도군'), ('SL_B', '용도'), ('SL_C', '시도'), ('SL_D', '시군구'), ('SL_E', '법원'),
           ('SL_F', '감정가 규모대'), ('SL_G', '남은 날'), ('SL_H', '조회수 구간')]
EMU, TOP, W, H, LEFT0, GAP = 12700, 48, 144, 170, 620, 160


def col_letter(n):
    s = ""
    while n: n, r = divmod(n - 1, 26); s = chr(65 + r) + s
    return s


def stage1(a):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    rows = list(csv.DictReader(open(a.csv, encoding="utf-8-sig")))
    recs = []
    for r in rows:
        d = dict(r)
        for k in ('감정가', '최저가'): d[k] = int(d[k]) if d[k] else None
        for k in ('최저가율', '건물면적', '토지면적'): d[k] = float(d[k]) if d[k] else None
        d['매각기일'] = date.fromisoformat(d['매각기일']) if d['매각기일'] else None
        d['조회수'] = None
        recs.append(d)
    N, NC = len(recs), len(HEAD)
    RAWHDR = 500 if N <= 420 else ((N + 80) // 100 + 1) * 100      # 02단원 [3] 원본 표 위치 규칙
    RAWDATA, RESEND = RAWHDR + 1, RAWHDR - 5
    LASTCOL, LASTROW = col_letter(NC), RAWHDR + N
    C = {h: i + 1 for i, (h, _) in enumerate(HEAD)}
    MONEY, DATES = {C['감정가'], C['최저가']}, {C['매각기일']}
    hf, hfill, base = Font(name=FONT, bold=True, color='FFFFFF', size=10), PatternFill('solid', fgColor='1F4E79'), Font(name=FONT, size=10)

    wb = Workbook(); ws = wb.active; ws.title = '탐색'
    m8 = re.search(r"(\d{8})", a.csv)
    ymd8 = m8.group(1) if m8 else date.today().strftime("%Y%m%d")
    title = a.title or ('신건레이더 %s — 법원경매정보(무료) · 전체 용도군 신건 %d건 (조회수는 법원경매정보 미제공)' % (ymd8, N))
    ws['B1'] = title; ws['B1'].font = Font(name=FONT, bold=True, size=14, color='1F4E79')
    ws['B2'] = '▼ 위 슬라이서를 눌러 조건을 좁히십시오. 선택 결과는 아래 15행부터 전부 표시됩니다. (해제: 슬라이서 우측 상단 지우개 아이콘)'
    ws['B2'].font = Font(name=FONT, size=9, color='595959')
    LIVE = [('선택 건수', '#,##0'), ('감정가 합계(원)', '#,##0'), ('감정가 평균(원)', '#,##0'),
            ('감정가 최대(원)', '#,##0'), ('최저가 합계(원)', '#,##0'), ('전체 대비 비율(%)', '0.0"%"')]
    for i, (lab, fmt) in enumerate(LIVE):
        r = 4 + i
        c = ws.cell(r, 2, lab); c.font = Font(name=FONT, bold=True, size=10, color='FFFFFF'); c.fill = hfill
        c.alignment = Alignment(horizontal='left', vertical='center')
        v = ws.cell(r, 3); v.font = Font(name=FONT, bold=True, size=13, color='C00000')
        v.fill = PatternFill('solid', fgColor='FCE4E4'); v.number_format = fmt
        v.alignment = Alignment(horizontal='right', vertical='center')
    ws['B13'] = '▼ 여기가 선택 결과입니다 (선택 조건에 맞는 전체 건수가 표시됩니다)'
    ws['B13'].font = Font(name=FONT, bold=True, size=10, color='C00000')

    def fmt(cell, j):
        cell.font = base
        if j in MONEY: cell.number_format = '#,##0'
        elif j in DATES: cell.number_format = 'yyyy-mm-dd'
        elif j == C['연번']: cell.number_format = '0'
        elif j == C['최저가율']: cell.number_format = '0.0'

    for j, (h, w) in enumerate(HEAD, start=1):
        c = ws.cell(TOPHDR, j, h); c.font = hf; c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[col_letter(j)].width = w
    ws.column_dimensions['B'].width = max(ws.column_dimensions['B'].width or 0, 18)   # 요약 KPI 가 잘리지 않게
    ws.column_dimensions['C'].width = max(ws.column_dimensions['C'].width or 0, 26)
    for r in range(15, RESEND + 1):
        for j in range(1, NC + 1): fmt(ws.cell(r, j), j)
    ws.cell(RAWHDR - 1, 2, '▼ 아래는 원본 표(슬라이서 필터 대상) — 보지 않으셔도 됩니다. 지우거나 숨기지 마십시오.').font = Font(name=FONT, bold=True, size=10, color='808080')
    for j, (h, _) in enumerate(HEAD, start=1):
        c = ws.cell(RAWHDR, j, h); c.font = hf; c.fill = hfill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for n, rec in enumerate(recs):
        rr = RAWDATA + n; rec['연번'] = n + 1
        for j, (k, _) in enumerate(HEAD, start=1):
            cell = ws.cell(rr, j, rec.get(k)); fmt(cell, j)
            if k == '사건번호': cell.number_format = '@'
    ws.cell(RAWHDR, 26, '표시여부(보조)').font = Font(name=FONT, bold=True, size=9, color='808080')
    for n in range(N):                                    # 보조열 Z — 빈 값 없는 키 열(사건번호 B) 참조
        ws.cell(RAWDATA + n, 26, '=SUBTOTAL(103,B%d)' % (RAWDATA + n)).font = Font(name=FONT, size=9)
    ws.column_dimensions['Z'].width = 13
    ws.row_dimensions[1].height = 22                      # 슬라이서(48+170pt)가 13·14행을 덮지 않게
    for r in range(4, 10): ws.row_dimensions[r].height = 24
    for r in range(10, 13): ws.row_dimensions[r].height = 20
    ws.freeze_panes = 'A15'; ws.sheet_view.showGridLines = False
    wb.create_sheet('피벗집계'); wb.create_sheet('대시보드'); wb.save(a.out)

    plan = {
        "file": a.out, "rows": N, "table": TABLE, "table_range": f"A{RAWHDR}:{LASTCOL}{LASTROW}",
        "kpi_formulas_C4_C9": [
            f"=SUBTOTAL(103,{TABLE}[사건번호])", f"=SUBTOTAL(109,{TABLE}[감정가])",
            f"=IFERROR(ROUND(SUBTOTAL(101,{TABLE}[감정가]),0),0)", f"=IFERROR(SUBTOTAL(104,{TABLE}[감정가]),0)",
            f"=SUBTOTAL(109,{TABLE}[최저가])",
            f"=IFERROR(ROUND(SUBTOTAL(103,{TABLE}[사건번호])/COUNTA({TABLE}[사건번호])*100,1),0)"],
        "spill_formula_A15": f'=LET(a,FILTER($A${RAWDATA}:${LASTCOL}${LASTROW},$Z${RAWDATA}:$Z${LASTROW}=1,""),IF(a="","",a))',
        "slicers": [{"name": n, "column": c, "position": f"{col_letter(8 + 2 * i)}3"} for i, (n, c) in enumerate(SLICERS)],
    }
    print("STAGE1 OK — 아래 값을 Excel MCP 2단계에 그대로 쓴다")
    print(json.dumps(plan, ensure_ascii=False, indent=1))


def patch(a):
    src = a.xlsx; tmp = src + ".patch.tmp"
    POS = {n: LEFT0 + GAP * i for i, (n, _) in enumerate(SLICERS)}
    CAP = dict(SLICERS)
    zin = zipfile.ZipFile(src); zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        d = zin.read(it.filename)
        if re.match(r'xl/drawings/drawing\d+\.xml$', it.filename):
            x = d.decode('utf-8')
            def repl(m):
                blk = m.group(0); nm = re.search(r'name="(SL_[^"]+)"', blk)
                if not nm or nm.group(1) not in POS: return blk
                head = ('<xdr:absoluteAnchor><xdr:pos x="%d" y="%d"/><xdr:ext cx="%d" cy="%d"/>'
                        % (POS[nm.group(1)] * EMU, TOP * EMU, W * EMU, H * EMU))
                inner = re.sub(r'^<xdr:twoCellAnchor[^>]*>', '', blk)
                inner = re.sub(r'<xdr:from>.*?</xdr:from>', '', inner, flags=re.S)
                inner = re.sub(r'<xdr:to>.*?</xdr:to>', '', inner, flags=re.S)
                inner = re.sub(r'</xdr:twoCellAnchor>$', '', inner)
                return head + inner + '</xdr:absoluteAnchor>'
            d = re.sub(r'<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>', repl, x, flags=re.S).encode('utf-8')
        elif re.match(r'xl/slicers/slicer\d+\.xml$', it.filename):
            x = d.decode('utf-8')
            for n, c in CAP.items():                      # 머리글만 열 이름으로. name 은 그대로(캐시 참조)
                x = re.sub(r'(<slicer\b[^>]*\bname="%s"[^>]*\bcaption=")[^"]*(")' % n, r'\g<1>%s\2' % c, x)
            d = x.encode('utf-8')
        zout.writestr(it, d)
    zout.close(); zin.close()

    za, zb = zipfile.ZipFile(src), zipfile.ZipFile(tmp)
    dr = zb.read('xl/drawings/drawing1.xml').decode()
    n_sl = len([n for n in zb.namelist() if n.startswith('xl/slicerCaches/')])
    checks = {
        "부품 수 동일": len(za.namelist()) == len(zb.namelist()),
        "twoCellAnchor 잔여 0": dr.count('twoCellAnchor') == 0,
        f"cx/cy 144×170pt {n_sl}개": dr.count('cx="1828800" cy="2159000"') == n_sl,
        "슬라이서 8개": n_sl == 8,
        "Requires sle15 보존": sum(zb.read(n).decode().count('Requires="sle15"') for n in zb.namelist() if n.endswith('.xml')) == n_sl,
    }
    for k, v in checks.items(): print(("PASS " if v else "FAIL ") + k)
    za.close(); zb.close()      # Windows 는 열린 zip 을 지우지 못한다 [확정 — 2026-09-30 Code 실측: WinError 32 → exit 1]
    ok = all(checks.values())
    if ok:
        with open(tmp, 'rb') as fi, open(src, 'wb') as fo: fo.write(fi.read())
        print("PATCH APPLIED →", src)
    else:
        print("검사 실패 — 원본 유지. 2단계(표·슬라이서)가 끝났는지 확인")
    try: os.remove(tmp)
    except OSError as e: print("임시 파일 삭제 실패(결과 파일은 정상):", tmp, e)
    if not ok: sys.exit(1)


def main():
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    p1 = sp.add_parser("stage1"); p1.add_argument("csv"); p1.add_argument("--out", required=True); p1.add_argument("--title")
    p2 = sp.add_parser("patch"); p2.add_argument("xlsx")
    a = ap.parse_args()
    (stage1 if a.cmd == "stage1" else patch)(a)


if __name__ == "__main__":
    main()
