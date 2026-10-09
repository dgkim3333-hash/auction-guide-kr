#!/usr/bin/env python3
"""
명도 일정관리 XLSX 빌더

4개 시트를 가진 일정관리 파일을 생성한다:
1. 마일스톤 - D+N 자동 계산 결과
2. 문서이력 - 생성·제출한 문서 기록
3. 커뮤니케이션 - 점유자와의 연락 로그
4. 비용트래커 - 송달료·이사비·집행비·소송비 누적

사용법:
    python build_schedule_xlsx.py \
      --case-no 2024타경12345 \
      --auction-date 2026-01-15 \
      --scenario A \
      --output /mnt/user-data/outputs/2024타경12345_명도일정관리.xlsx
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 같은 디렉토리의 calculate_milestones 사용
sys.path.insert(0, str(Path(__file__).parent))
from calculate_milestones import calculate as calc_milestones  # noqa: E402

# 표준 폰트 (xlsx 스킬 가이드 준수)
DEFAULT_FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name=DEFAULT_FONT, color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(name=DEFAULT_FONT, size=10)
INPUT_FONT = Font(name=DEFAULT_FONT, size=10, color="0000FF")  # 파란색 = 사용자 입력
FORMULA_FONT = Font(name=DEFAULT_FONT, size=10, color="000000")  # 검정 = 수식
WARN_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


SCENARIO_NAMES = {
    "A": "일반 채무자/소유자 점유",
    "B": "명도합의 우선",
    "C": "대항력 있는 임차인",
    "D": "배당받은 임차인",
    "E": "유치권 주장",
}


def style_header_row(ws, row_idx: int, cols: int):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize(ws, min_w: int = 10, max_w: int = 50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            longest = max(longest, max(len(line) for line in v.splitlines()) if v else 0)
        ws.column_dimensions[col_letter].width = max(min_w, min(longest + 3, max_w))


def build_milestone_sheet(wb: Workbook, case_no: str, auction_date_str: str, scenario: str):
    ws = wb.create_sheet("마일스톤", 0)
    d0 = datetime.strptime(auction_date_str, "%Y-%m-%d").date()

    # 헤더 정보 영역
    ws["A1"] = "사건번호"
    ws["B1"] = case_no
    ws["A2"] = "낙찰일 (D+0)"
    ws["B2"] = d0
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["A3"] = "점유자 시나리오"
    ws["B3"] = f"{scenario} - {SCENARIO_NAMES.get(scenario, '미정')}"
    for r in (1, 2, 3):
        ws.cell(row=r, column=1).font = Font(name=DEFAULT_FONT, bold=True)
        ws.cell(row=r, column=2).font = INPUT_FONT
        ws.cell(row=r, column=2).fill = WARN_FILL

    # 마일스톤 표 시작
    headers = ["D+N", "예상일자", "요일", "마일스톤", "카테고리", "근거", "상태", "메모"]
    header_row = 5
    for i, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=i, value=h)
    style_header_row(ws, header_row, len(headers))

    rows = calc_milestones(d0)
    for i, m in enumerate(rows, start=header_row + 1):
        ws.cell(row=i, column=1, value=f"D+{m['days_from_d0']}")
        ws.cell(row=i, column=2, value=datetime.strptime(m["date"], "%Y-%m-%d").date())
        ws.cell(row=i, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=3, value=m["weekday"])
        ws.cell(row=i, column=4, value=m["milestone"])
        ws.cell(row=i, column=5, value=m["category"])
        ws.cell(row=i, column=6, value=m["basis"])
        ws.cell(row=i, column=7, value="대기")  # 사용자가 진행/완료로 변경
        ws.cell(row=i, column=8, value="")

        # 카테고리별 색상 미세 차이
        if m["category"] == "한도":
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
        for c in range(1, len(headers) + 1):
            ws.cell(row=i, column=c).font = BODY_FONT
            ws.cell(row=i, column=c).border = BORDER
            ws.cell(row=i, column=c).alignment = LEFT_WRAP if c in (4, 6, 8) else CENTER

    # 사용자 안내
    note_row = header_row + len(rows) + 2
    ws.cell(row=note_row, column=1, value="⚠️ 안내")
    ws.cell(row=note_row, column=1).font = Font(name=DEFAULT_FONT, bold=True, color="C00000")
    ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=8)
    ws.cell(
        row=note_row,
        column=2,
        value="법정 기한은 자연일 기준. 공휴일 만료/송달지연/항고 발생 시 변동. "
              "인지액·송달료는 대법원 전자소송에서 최신 기준 확인 필수.",
    )
    ws.cell(row=note_row, column=2).alignment = LEFT_WRAP
    ws.cell(row=note_row, column=2).font = Font(name=DEFAULT_FONT, size=10, italic=True)

    autosize(ws)
    ws.freeze_panes = "A6"


def build_documents_sheet(wb: Workbook):
    ws = wb.create_sheet("문서이력")
    headers = ["순번", "작성일", "문서명", "시나리오", "수신인/제출처", "송달방법", "상태", "비고"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    # 빈 행 20개 (사용자가 채움)
    for r in range(2, 22):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = INPUT_FONT
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).alignment = LEFT_WRAP if c in (3, 5, 8) else CENTER
    ws.cell(row=2, column=1, value=1)
    autosize(ws)
    ws.freeze_panes = "A2"


def build_communication_sheet(wb: Workbook):
    ws = wb.create_sheet("커뮤니케이션")
    headers = ["순번", "일자", "수단", "상대방", "주요내용", "다음액션", "후속기한"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    for r in range(2, 32):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = INPUT_FONT
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).alignment = LEFT_WRAP if c in (5, 6) else CENTER
    ws.cell(row=2, column=1, value=1)
    autosize(ws)
    ws.freeze_panes = "A2"


def build_cost_tracker_sheet(wb: Workbook):
    ws = wb.create_sheet("비용트래커")
    headers = ["순번", "일자", "항목", "분류", "금액(원)", "지급처", "증빙", "비고"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))

    # 카테고리 예시 행 (사용자가 채울 수 있는 가이드)
    sample_rows = [
        ("인도명령 인지", "법원수수료", 1000),
        ("인도명령 송달료", "법원수수료", 50000),
        ("강제집행 예납금", "집행비용", None),
        ("이사비", "협상비용", None),
        ("변호사 선임료", "법률비용", None),
        ("소장 인지", "법원수수료", None),
    ]
    for r, (item, cat, amount) in enumerate(sample_rows, start=2):
        ws.cell(row=r, column=1, value=r - 1)
        ws.cell(row=r, column=3, value=item)
        ws.cell(row=r, column=4, value=cat)
        if amount is not None:
            ws.cell(row=r, column=5, value=amount)
            ws.cell(row=r, column=5).number_format = '#,##0;(#,##0);"-"'

    # 빈 행 추가
    last_data_row = 1 + len(sample_rows)
    for r in range(last_data_row + 1, last_data_row + 21):
        ws.cell(row=r, column=1, value=r - 1)

    # 합계 행 (수식)
    total_row = last_data_row + 22
    ws.cell(row=total_row, column=4, value="합계")
    ws.cell(row=total_row, column=4).font = Font(name=DEFAULT_FONT, bold=True)
    ws.cell(row=total_row, column=4).alignment = CENTER
    ws.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row - 1})")
    ws.cell(row=total_row, column=5).font = FORMULA_FONT
    ws.cell(row=total_row, column=5).number_format = '#,##0;(#,##0);"-"'
    ws.cell(row=total_row, column=5).fill = PatternFill("solid", fgColor="FFF2CC")

    # 모든 데이터 셀 스타일링
    for r in range(2, total_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER
            if r != total_row and c == 5:
                ws.cell(row=r, column=c).font = INPUT_FONT
                ws.cell(row=r, column=c).number_format = '#,##0;(#,##0);"-"'
            if c in (3, 6, 7, 8) and r != total_row:
                ws.cell(row=r, column=c).font = INPUT_FONT
                ws.cell(row=r, column=c).alignment = LEFT_WRAP
            elif c not in (5,) and r != total_row:
                ws.cell(row=r, column=c).alignment = CENTER

    autosize(ws)
    ws.freeze_panes = "A2"


def main():
    parser = argparse.ArgumentParser(description="명도 일정관리 XLSX 빌더")
    parser.add_argument("--case-no", required=True, help="사건번호 (예: 2024타경12345)")
    parser.add_argument("--auction-date", required=True, help="낙찰일 YYYY-MM-DD")
    parser.add_argument(
        "--scenario",
        required=True,
        choices=list(SCENARIO_NAMES.keys()),
        help="점유자 시나리오 (A/B/C/D/E)",
    )
    parser.add_argument("--output", required=True, help="출력 XLSX 경로")
    args = parser.parse_args()

    wb = Workbook()
    # 기본 시트 제거
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    build_milestone_sheet(wb, args.case_no, args.auction_date, args.scenario)
    build_documents_sheet(wb)
    build_communication_sheet(wb)
    build_cost_tracker_sheet(wb)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"✅ 생성 완료: {output_path}")
    print(f"   사건번호: {args.case_no}")
    print(f"   낙찰일: {args.auction_date}")
    print(f"   시나리오: {args.scenario} - {SCENARIO_NAMES[args.scenario]}")


if __name__ == "__main__":
    main()
