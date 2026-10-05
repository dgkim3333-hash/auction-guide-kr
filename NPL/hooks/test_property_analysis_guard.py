# -*- coding: utf-8 -*-
# property_analysis_guard 변이 시험 — 막아야 할 때 막고, 통과시킬 때 통과시키는지
import json, os, subprocess, sys, tempfile

G = os.path.join(os.path.dirname(os.path.abspath(__file__)), "property_analysis_guard.py")


def run(mode, payload):
    r = subprocess.run([sys.executable, G, mode], input=json.dumps(payload, ensure_ascii=False).encode("utf-8"), capture_output=True)
    return r.returncode, r.stdout.decode("utf-8", "replace"), r.stderr.decode("utf-8", "replace")


def tr(events):
    f = tempfile.NamedTemporaryFile("w", suffix=".jsonl", delete=False, encoding="utf-8")
    for e in events:
        f.write(json.dumps(e, ensure_ascii=False) + "\n")
    f.close()
    return f.name


U = lambda t: {"type": "user", "message": {"role": "user", "content": t}}
SK = lambda s: {"type": "assistant", "message": {"role": "assistant", "content": [{"type": "tool_use", "name": "Skill", "input": {"skill": s}}]}}
BUILD = {"type": "assistant", "message": {"role": "assistant", "content": [{"type": "tool_use", "name": "PowerShell", "input": {"command": "node _card.js model.json 03_산출물\\물건카드_2025타경1.pptx"}}]}}
RES = lambda t: {"type": "user", "message": {"role": "user", "content": [{"type": "tool_result", "content": t}]}}

cases = []
# 1차: 요청 감지
rc, out, _ = run("prompt", {"prompt": "2025타경12345 이 물건 분석해줘"}); cases.append(("prompt 감지 → 체크리스트 주입", rc == 0 and "필수 절차" in out))
rc, out, _ = run("prompt", {"prompt": "오늘 날씨 어때"}); cases.append(("prompt 무관 → 주입 없음", rc == 0 and out.strip() == ""))
# 2차: 종료 검사
base = [U("2025타경12345 분석해줘"), SK("anthropic-skills:npl-analysis"), SK("anthropic-skills:auction-property-card"), BUILD]
rc, _, err = run("stop", {"transcript_path": tr(base)}); cases.append(("산출물 생성 + 검증 없음 → 차단", rc == 2 and "검증 미통과" in err))
rc, _, _ = run("stop", {"transcript_path": tr(base + [RES("...\nVERIFY_PROPERTY PASS")])}); cases.append(("검증 PASS → 통과", rc == 0))
rc, _, _ = run("stop", {"transcript_path": tr(base + [RES("VERIFY_PROPERTY FAIL")])}); cases.append(("검증 FAIL → 차단", rc == 2))
noskill = [U("2025타경12345 분석해줘"), SK("anthropic-skills:npl-analysis"), BUILD, RES("VERIFY_PROPERTY PASS")]
rc, _, err = run("stop", {"transcript_path": tr(noskill)}); cases.append(("auction-property-card 미호출 → 차단", rc == 2 and "auction-property-card" in err))
rc, out, _ = run("stop", {"transcript_path": tr(base), "stop_hook_active": True}); cases.append(("두 번째 종료 → 화면 경고 후 통과", rc == 0 and "systemMessage" in out))
rc, _, _ = run("stop", {"transcript_path": tr([U("종부세 개편안 설명해줘")])}); cases.append(("분석 아닌 차례 → 통과", rc == 0))


# ── 2026-10-05 추가: 오탐 방지 · 관점 2 다섯 장 검사 ──
W = lambda path, content: {"type": "assistant", "message": {"role": "assistant", "content": [{"type": "tool_use", "name": "Write", "input": {"file_path": path, "content": content}}]}}
SH = lambda cmd: {"type": "assistant", "message": {"role": "assistant", "content": [{"type": "tool_use", "name": "PowerShell", "input": {"command": cmd}}]}}
peek = [U("두 덱 비교해줘"), W(r"C:\tmp\deck_peek.py", 'NEW = r"C:\\AI\\NPL\\x\\03_산출물\\물건카드_2025타경1.pptx"'), SH(r"python C:\tmp\deck_peek.py")]
rc, _, _ = run("stop", {"transcript_path": tr(peek)}); cases.append(("읽기 스크립트(본문에 pptx 경로) → 통과(오탐 없음)", rc == 0))
editcard = [U("물건 정리"), W(r"C:\AI\NPL\x\_card.js", "// 생성기"), SH("node _card.js")]
rc, _, _ = run("stop", {"transcript_path": tr(editcard)}); cases.append(("_card.js 저장·실행 → 생성으로 판정 → 차단", rc == 2))
try:
    import shutil as _sh
    from pptx import Presentation
    from pptx.util import Inches
    V = os.path.join(os.path.dirname(os.path.abspath(__file__)), "verify_property.py")
    def mk(folder, text):
        os.makedirs(os.path.join(folder, "01_원본서류")); os.makedirs(os.path.join(folder, "03_산출물"))
        for n in ["매각물건명세서.pdf", "건물등기.pdf", "토지등기.pdf", "건축물대장.pdf", "감정평가서.pdf"]:
            open(os.path.join(folder, "01_원본서류", n), "wb").close()
        pr = Presentation(); sl = pr.slides.add_slide(pr.slide_layouts[5])
        sl.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(4)).text_frame.text = text
        sl.notes_slide.notes_text_frame.text = "노트"; pr.save(os.path.join(folder, "03_산출물", "물건카드_2025타경1.pptx"))
    root = tempfile.mkdtemp()
    a, b = os.path.join(root, "a"), os.path.join(root, "b")
    mk(a, "배당 순서"); mk(b, "배당 순서 · 수요처 · 확인 방법 · 인도명령 · 배점")
    oa = subprocess.run([sys.executable, V, a], capture_output=True).stdout.decode("utf-8", "replace")
    ob = subprocess.run([sys.executable, V, b], capture_output=True).stdout.decode("utf-8", "replace")
    cases.append(("verify 다섯 장 없는 덱 → 4개 FAIL", all(f"[FAIL] {k}" in oa for k in ["명도", "종합점수", "수요처", "고유리스크"])))
    cases.append(("verify 다섯 장 있는 덱 → 4개 PASS", all(f"[PASS] {k}" in ob for k in ["명도", "종합점수", "수요처", "고유리스크"])))
    _sh.rmtree(root, ignore_errors=True)
except ImportError:
    cases.append(("python-pptx 없음 — 검증기 시험 생략", False))

ok = all(c[1] for c in cases)
for n, r in cases:
    print(("PASS " if r else "FAIL ") + n)
print("ALL PASS" if ok else "SOME FAIL")
sys.exit(0 if ok else 1)
