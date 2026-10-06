# -*- coding: utf-8 -*-
"""물건 분석 산출물 검증기 — 수강생판 (강사 원본 2026-10-05, 「빠짐없이 분석」 2차 장치)

사용:  python verify_property.py "<물건 폴더>"
결과:  마지막 줄에 VERIFY_PROPERTY PASS / VERIFY_PROPERTY FAIL
예외:  물건 폴더의 _검증예외.md 에 「- <검사ID>: 사유」 줄이 있으면 그 검사는 WAIVED(사유 필수)

필요: pip install python-pptx openpyxl
검사 기준 = C:\\AI\\NPL\\CLAUDE.md ■ 이원 분석·두 관점·세금·대출·산출물 원칙
           + npl-analysis / auction-property-card / npl-excel-fill-map 스킬 필수 항목 + auction-failure-guard 실패위험 점검표
"""
import glob, os, re, sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# (검사ID, 설명, 정규식 목록 — 모두 만족해야 통과)
DECK_CHECKS = [
    ("두관점", "관점 1(NPL 매입 후 제3자 낙찰)·관점 2(직접 낙찰) 둘 다", [r"관점\s*1", r"관점\s*2"]),
    ("3열비교", "배당 회수 / 유입(개인) / 유입(법인) 3열 비교", [r"배당\s*회수", r"유입.{0,3}개인", r"유입.{0,3}법인"]),
    ("취득세", "취득세 개인·법인", [r"취득세"]),
    ("보유세", "재산세·종부세(첫 과세연도)", [r"재산세", r"종부세"]),
    ("출구세", "출구세(산출 또는 「출구 시점 지정 필요」)", [r"출구"]),
    ("선순위조세", "STEP 4-C 선순위 조세 X·법정기일", [r"법정기일|선순위\s*조세"]),
    ("당해세", "당해세 판정", [r"당해세"]),
    ("최우선", "소액임차인 최우선변제 판정", [r"최우선"]),
    ("정비사업", "서울 정비사업 인접 체크(경기는 「데이터 없음」 명시)", [r"정비"]),
    ("역세권", "역세권 실측(STEP 5-B)", [r"역세권"]),
    ("추정감정가", "추정 감정가 표(원칙 24)", [r"추정\s*감정가"]),
    ("시세출처", "시세 출처 게이트(창고 수집일·조회일 또는 범위 밖 사유)", [r"(창고|DuckDB).{0,20}(수집|적재)|(창고|DuckDB)\s*범위\s*밖|조회일"]),
    ("임대", "임대 가정 — 전월세 실거래 중앙값(원칙 25)", [r"월세", r"전월세"]),
    ("레버리지", "대출 판정·Cash-on-Cash", [r"Cash-on-Cash|현금수익률", r"대출"]),
    ("금리", "ECOS(한국은행) 금리 근거", [r"ECOS|한국은행"]),
    ("체크리스트", "21개 체크리스트", [r"체크리스트"]),
    ("문건", "Agent 0 문건 분석·저항지수", [r"문건", r"저항"]),
    ("선점", "채권양수 선점 체크", [r"선점"]),
    ("민감도", "민감도(감정가=추정·X·지연 등)", [r"민감도"]),
    ("확인사항", "확인해야 할 것(우선순위)", [r"확인"]),
    ("실패점검", "실패위험 점검표(auction-failure-guard 11개 유형)", [r"실패\s*위험\s*점검표"]),
    # 관점 2 다섯 장 (2026-10-05)
    ("수요처", "입지 장 — 주요 수요처(대학·산업단지·업무지구) 직선거리", [r"수요처"]),
    ("고유리스크", "고유 리스크 장 — 「리스크 / 사실 / 영향 / 확인 방법」 표", [r"확인\s*방법"]),
    ("명도", "명도 장 — 인도명령 일정·예산", [r"인도명령"]),
    ("종합점수", "종합 점수 장 — 100점 배점표(관점 2)", [r"배점"]),
]
ABBR = re.compile(r"\d[\d,.]*\s*(?:억|만\s*원|조\s*원)")


def deck_text(pptx_path):
    from pptx import Presentation
    p = Presentation(pptx_path)
    parts, pics, no_notes = [], 0, []
    for i, s in enumerate(p.slides, 1):
        for sh in s.shapes:
            if sh.shape_type == 13:
                pics += 1
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        parts.append(c.text)
        if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
            no_notes.append(i)
    return "\n".join(parts), pics, no_notes, len(p.slides)


def excel_check(xlsx):
    import openpyxl
    REQ = ['C3', 'C4', 'C5', 'E5', 'C6', 'C7', 'C8', 'C9', 'C10', 'C11', 'C12', 'C13', 'E6', 'E8', 'E9', 'E11', 'G3',
           'C31', 'C32', 'C33', 'C34', 'C35', 'C37', 'C39', 'C40', 'C41', 'C42', 'C43', 'C44', 'C45',
           'F30', 'F31', 'F32', 'F33', 'F34', 'F35']
    ws = openpyxl.load_workbook(xlsx)['분석']
    bad = [c for c in REQ if ws[c].value in (None, '') and ws[c].comment is None]
    if str(ws['C36'].value).replace(' ', '') != '=SUM(C39:C44)':
        bad.append('C36 수식')
    return bad


def main(folder):
    folder = os.path.abspath(folder)
    waive = {}
    wf = os.path.join(folder, "_검증예외.md")
    if os.path.exists(wf):
        for line in open(wf, encoding="utf-8"):
            m = re.match(r"\s*-\s*([^:：]+)[:：]\s*(.+)", line)
            if m and m.group(2).strip():
                waive[m.group(1).strip()] = m.group(2).strip()
    res = []   # (id, ok, detail)

    def rec(cid, ok, detail=""):
        if not ok and cid in waive:
            res.append((cid, "WAIVED", f"{detail} — 예외 사유: {waive[cid]}"))
        else:
            res.append((cid, "PASS" if ok else "FAIL", detail))

    # 1. 폴더·서류
    src = os.path.join(folder, "01_원본서류")
    names = " ".join(os.path.basename(p) for p in glob.glob(os.path.join(src, "*")))
    rec("서류_명세서", bool(re.search(r"매각물건명세서|공매재산명세서", names)), "01_원본서류 매각물건명세서(공매재산명세서)")
    rec("서류_등기", bool(re.search(r"등기", names)), "01_원본서류 등기부(건물·토지)")
    rec("서류_토지등기", bool(re.search(r"토지", names)) or "집합" in names, "토지 등기부(집합건물이면 예외 사유 기재)")
    rec("서류_대장", bool(re.search(r"건축물대장", names)), "01_원본서류 건축물대장")
    rec("서류_감정", bool(re.search(r"감정평가", names)), "01_원본서류 감정평가서(예정물건이면 예외 사유)")
    card = os.path.join(folder, "00_물건카드.md")
    ok = os.path.exists(card) and re.search(r"\d{4}\s*타경\s*\d+|\d{4}-\d{5}-\d{3}", open(card, encoding="utf-8").read()) is not None
    rec("물건카드md", bool(ok), "00_물건카드.md 작성(템플릿 그대로가 아님)")

    # 2. 엑셀 2종
    npl = sorted(glob.glob(os.path.join(folder, "02_분석", "NPL수익률_*.xlsx")), key=os.path.getmtime)
    rec("NPL수익률엑셀", bool(npl), "02_분석\\NPL수익률_<사건번호>_<주소>.xlsx")
    if npl:
        bad = excel_check(npl[-1])
        rec("NPL수익률_빈칸", not bad, f"값·메모 없는 칸: {bad}" if bad else "정본표 빈칸 검사 통과")
    ana = [p for p in glob.glob(os.path.join(folder, "02_분석", "*.xlsx")) if "NPL수익률" not in os.path.basename(p) and not os.path.basename(p).startswith("~$")]
    rec("분석엑셀", bool(ana), "02_분석 셀 수식 분석 엑셀(안분·세금·레버리지)")

    # 3. PPTX
    decks = sorted([p for p in glob.glob(os.path.join(folder, "03_산출물", "*.pptx")) if not os.path.basename(p).startswith("~$")], key=os.path.getmtime)
    rec("PPTX", bool(decks), "03_산출물 물건카드 PPTX")
    if decks:
        text, pics, no_notes, n = deck_text(decks[-1])
        for cid, desc, pats in DECK_CHECKS:
            miss = [p for p in pats if not re.search(p, text)]
            rec(cid, not miss, desc + (f" — 없는 표현: {miss}" if miss else ""))
        rec("도표3종", pics >= 3, f"그림 {pics}개(권리 사다리·문건 타임라인·배당 폭포 ≥3)")
        ab = ABBR.findall(text)
        rec("금액표기", not ab, f"억·만 축약 {ab[:5]}" if ab else "원 단위 풀 콤마")
        rec("발표자노트", not no_notes, f"노트 없는 장 {no_notes}" if no_notes else f"{n}장 전부 노트")
        tax_slides = len(re.findall(r"세금[①②③④]|세금\s*\d", text))
        rec("세금장3장", tax_slides >= 3, f"세금 장 표기 {tax_slides}개(최소 3장)")

    fails = [r for r in res if r[1] == "FAIL"]
    for cid, st, d in res:
        print(f"[{st}] {cid} — {d}")
    print(f"검사 {len(res)} · 실패 {len(fails)} · 예외 {sum(1 for r in res if r[1] == 'WAIVED')} · 폴더 {folder}")
    print("VERIFY_PROPERTY PASS" if not fails else "VERIFY_PROPERTY FAIL")
    return 0 if not fails else 1


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("사용: python verify_property.py <물건 폴더>")
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
